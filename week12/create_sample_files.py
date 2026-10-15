# agentic-comm-automator/create_sample_files.py
# สร้างไฟล์แนบตัวอย่าง data/reports/sales_report.pdf และ data/reports/summary_data.xlsx
import os
import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORTS = os.path.join(BASE_DIR, "data", "reports")

SALES = [("Alice", 1425.75), ("Bob", 130.75), ("Charlie", 500.00), ("David", 90.00), ("Eve", 15.50)]


def create_pdf(path):
    c = canvas.Canvas(path, pagesize=A4)
    w, h = A4
    c.setFillColor(HexColor("#1F4E79"))
    c.rect(0, h - 90, w, 90, fill=1, stroke=0)
    c.setFillColor(HexColor("#FFFFFF"))
    c.setFont("Helvetica-Bold", 22)
    c.drawString(50, h - 55, "Daily Sales Report")
    c.setFont("Helvetica", 11)
    c.drawString(50, h - 75, f"Generated {datetime.date.today():%Y-%m-%d}")

    y = h - 140
    c.setFillColor(HexColor("#000000"))
    c.setFont("Helvetica-Bold", 12)
    c.drawString(60, y, "Customer")
    c.drawRightString(400, y, "Total Amount ($)")
    c.setFont("Helvetica", 12)
    for name, amount in SALES:
        y -= 22
        c.drawString(60, y, name)
        c.drawRightString(400, y, f"{amount:,.2f}")
    y -= 28
    c.setFont("Helvetica-Bold", 12)
    c.drawString(60, y, "Total")
    c.drawRightString(400, y, f"{sum(a for _, a in SALES):,.2f}")
    c.save()
    print(f"Created {path}")


def create_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Customer", "Total_Amount"])
    for row in SALES:
        ws.append(row)
    ws.append(["Total", f"=SUM(B2:B{len(SALES) + 1})"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    wb.save(path)
    print(f"Created {path}")


if __name__ == "__main__":
    os.makedirs(REPORTS, exist_ok=True)
    create_pdf(os.path.join(REPORTS, "sales_report.pdf"))
    create_xlsx(os.path.join(REPORTS, "summary_data.xlsx"))
